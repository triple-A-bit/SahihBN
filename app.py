import streamlit as st
import pytesseract
from PIL import Image
import pandas as pd
from docx import Document
import io

# --- CONFIGURATION ---
st.set_page_config(page_title="Halal Verify OCR", page_icon="✅")

# --- FUNCTIONS ---

def extract_details(ocr_text):
    """
    Parses OCR text to find key details based on keywords.
    Note: This is basic keyword matching. For an FYP, you can claim this is 
    'Rule-based Information Extraction'.
    """
    lines = ocr_text.split('\n')
    data = {
        "Product Name": "",
        "Ingredients": "",
        "Manufacturer": "",
        "Halal Certified": False
    }

    # 1. Simple Halal Detection (Keyword Search)
    halal_keywords = ["halal", "ms 1500", "muib", "jakim", "brunei"]
    if any(keyword in ocr_text.lower() for keyword in halal_keywords):
        data["Halal Certified"] = True

    # 2. Logic to find fields (Heuristic)
    # We look for lines containing "Ingredients" or "Manufactured by" 
    # and take the text following them.
    for i, line in enumerate(lines):
        line_lower = line.lower()
        
        # Product Name Heuristic: Usually the first non-empty line with large text
        # (Hard to detect size here, so we default to the first meaningful line if empty)
        if not data["Product Name"] and len(line) > 3:
            data["Product Name"] = line.strip()

        # Ingredients Extraction
        if "ingredients" in line_lower or "ramuan" in line_lower:
            # Get the current line plus next few lines until a blank line
            ingredients_text = line.split(":", 1)[-1].strip()
            # If the ingredient list continues to next lines
            for j in range(i + 1, len(lines)):
                if lines[j].strip() == "": break
                ingredients_text += " " + lines[j].strip()
            data["Ingredients"] = ingredients_text

        # Manufacturer Extraction
        if "manufactured by" in line_lower or "dibuat oleh" in line_lower:
            data["Manufacturer"] = line.split(":", 1)[-1].strip()
            if not data["Manufacturer"] and i+1 < len(lines):
                data["Manufacturer"] = lines[i+1].strip()

    return data

def to_excel(df):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Sheet1')
    return output.getvalue()

def to_word(data_dict):
    doc = Document()
    doc.add_heading('Halal Verification Report', 0)
    
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Category'
    hdr_cells[1].text = 'Details'

    for key, value in data_dict.items():
        row_cells = table.add_row().cells
        row_cells[0].text = str(key)
        row_cells[1].text = str(value)
        
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()

# --- APP UI ---

st.title("🔍 Halal Product Scanner")
st.write("Upload a product image to extract details automatically.")

uploaded_file = st.file_uploader("Choose an image...", type=["jpg", "png", "jpeg"])

if uploaded_file is not None:
    # 1. Show Image
    image = Image.open(uploaded_file)
    st.image(image, caption='Uploaded Product', use_column_width=True)
    
    with st.spinner('Scanning text...'):
        # 2. Perform OCR
        try:
            # In Termux, tesseract is usually in path. If not, uncomment line below:
            # pytesseract.pytesseract.tesseract_cmd = '/data/data/com.termux/files/usr/bin/tesseract'
            raw_text = pytesseract.image_to_string(image)
        except Exception as e:
            st.error(f"OCR Error: {e}. Did you install tesseract?")
            raw_text = ""

    if raw_text:
        # 3. Extract Data
        extracted_data = extract_details(raw_text)

        st.subheader("📝 Verify & Edit Data")
        
        # 4. Form for Editing
        with st.form("edit_form"):
            col1, col2 = st.columns(2)
            
            with col1:
                new_name = st.text_input("Product Name", extracted_data["Product Name"])
                new_manufacturer = st.text_input("Manufactured By", extracted_data["Manufacturer"])
            
            with col2:
                # Checkbox for Halal
                new_halal = st.checkbox("Halal Certified (Tick if logo is present)", value=extracted_data["Halal Certified"])
            
            new_ingredients = st.text_area("Ingredients", extracted_data["Ingredients"], height=150)
            
            submitted = st.form_submit_button("Confirm Data")

        if submitted:
            final_data = {
                "Product Name": new_name,
                "Manufacturer": new_manufacturer,
                "Ingredients": new_ingredients,
                "Halal Certified": "Yes" if new_halal else "No"
            }
            
            # Display Final Table
            st.success("Data Confirmed!")
            df = pd.DataFrame([final_data])
            st.table(df)

            # 5. Download Options
            col_d1, col_d2 = st.columns(2)
            
            with col_d1:
                st.download_button(
                    label="📥 Download Excel",
                    data=to_excel(df),
                    file_name='product_data.xlsx',
                    mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
                )
            
            with col_d2:
                st.download_button(
                    label="📥 Download Word",
                    data=to_word(final_data),
                    file_name='product_data.docx',
                    mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
                )

    with st.expander("See raw OCR text"):
        st.text(raw_text)
